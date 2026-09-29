"""Structural, numerical and rendered-page checks for the generated paper."""
from pathlib import Path
from collections import Counter
import hashlib, json, re, sys
import pandas as pd
import numpy as np
import fitz
from PIL import Image, ImageOps, ImageDraw
from docx import Document
from lxml import etree, html

ROOT=Path(__file__).resolve().parents[1]
sys.stdout.reconfigure(encoding='utf-8')
OUT=ROOT/'output/working-paper'
QA=OUT/'qa';QA.mkdir(exist_ok=True)
STEM='ChatGPT_Croatia_Working_Paper'
doc=Document(OUT/(STEM+'.docx'))
pdf=fitz.open(OUT/(STEM+'.pdf'))
web=html.parse(str(OUT/(STEM+'.html')))
body=web.getroot().text_content()
md=(OUT/(STEM+'.md')).read_text(encoding='utf-8')
alltext='\n'.join(page.get_text() for page in pdf)
assert '{{' not in md
for bad in ['[AUTHOR','TODO','FIXME','Citation not found','Error! Reference source not found','�']:
    assert bad not in body and bad not in alltext, bad
assert len(web.findall('.//table'))==len(doc.tables)==14
assert len(web.findall('.//img'))==len(doc.inline_shapes)==3
assert len(doc._element.xpath('.//m:oMath'))>=4, 'Missing native Word math'
assert web.xpath('//math'), 'Missing HTML MathML'
references=web.xpath('//*[@id="refs"]/*[contains(@class,"csl-entry")]')
assert len(references)>=25
cited=set(re.findall(r'data-cites="([^"]+)"',(OUT/(STEM+'.html')).read_text(encoding='utf-8')))
refids={x.get('id').removeprefix('ref-') for x in references}
assert set(' '.join(cited).split()) <= refids
assert all(img.get('src','').startswith('data:') for img in web.findall('.//img'))

monthly=pd.read_csv(OUT/'analysis/monthly_all.csv')
assert np.allclose(monthly.skills_n,monthly.SKILLS*monthly.n/100)
prior=pd.read_csv(ROOT/'output/reviews/paper1_2026-09-29/monthly_diagnostics.csv')
assert np.array_equal(prior.n,monthly.n)
assert np.allclose(prior.threat,monthly.threat)
assert np.allclose(prior.skills,monthly.SKILLS)
duplicates=pd.read_csv(OUT/'analysis/duplicates.csv')
assert np.allclose(duplicates.share,duplicates.matched_n/duplicates.n*100)

page_info=[]
for i,page in enumerate(pdf):
    text=page.get_text()
    assert len(text.strip())>30, ('Empty page',i+1)
    pix=page.get_pixmap(matrix=fitz.Matrix(1.7,1.7),alpha=False)
    pix.save(QA/f'page-{i+1:02d}.png')
    overflow=[]
    for block in page.get_text('dict')['blocks']:
        if block['type']!=0:continue
        for line in block['lines']:
            for span in line['spans']:
                x0,y0,x1,y1=span['bbox']
                if x0<15 or x1>page.rect.width-15 or y0<12 or y1>page.rect.height-12:
                    overflow.append(span['text'])
    page_info.append({'page':i+1,'characters':len(text),'first_lines':text.splitlines()[:5],
        'last_lines':text.splitlines()[-5:],'edge_overflow':overflow})
    assert not overflow, (i+1,overflow)

for batch in range(0,len(pdf),6):
    sheet=Image.new('RGB',(1500,2020),'#dddddd');draw=ImageDraw.Draw(sheet)
    for k,j in enumerate(range(batch,min(batch+6,len(pdf)))):
        im=Image.open(QA/f'page-{j+1:02d}.png');im.thumbnail((485,950))
        x=(k%3)*500+(500-im.width)//2;y=(k//3)*1010+35
        sheet.paste(im,(x,y));draw.text((x,y-23),f'PAGE {j+1}',fill='black')
    sheet.save(QA/f'overview-{batch//6+1:02d}.png')

report={'pages':len(pdf),'tables':len(doc.tables),'figures':len(doc.inline_shapes),
    'references':len(references),'native_word_equations':len(doc._element.xpath('.//m:oMath')),
    'html_self_contained':True,'numerical_checks':'passed','page_checks':page_info,
    'rendering':'Native Microsoft Word PDF export; PyMuPDF page rasterisation. Packaged renderer unavailable: pdf2image and LibreOffice not installed.',
    'files':{p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        for p in [OUT/(STEM+e) for e in ['.pdf','.html','.docx']]}}
(QA/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'extracted_text.txt').write_text(alltext,encoding='utf-8')
print(json.dumps({k:report[k] for k in ['pages','tables','figures','references','native_word_equations','numerical_checks']}))
for r in page_info:print(r['page'],r['characters'],' | '.join(r['first_lines'][:2]),'...',' | '.join(r['last_lines'][-2:]))
