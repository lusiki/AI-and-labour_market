"""Build a single-source scholarly working paper (HTML and editable DOCX).

Run R/04_working_paper_analysis.R first. PDF export and visual QA are performed
by R/04_export_working_paper.ps1 and R/04_check_working_paper.py.
"""
from pathlib import Path
import csv, hashlib, json, re, subprocess
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/working-paper'
A = OUT / 'analysis'
ASSETS = OUT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
STEM = 'ChatGPT_Croatia_Working_Paper'
PANDOC = Path('C:/Program Files/Quarto/bin/tools/pandoc.exe')
D = {p.stem: pd.read_csv(p) for p in A.glob('*.csv')}
S = D['summary'].iloc[0]
pre, post = [D['period_monthly'].query('post == @v').iloc[0] for v in [0, 1]]
bp, ba = [D['balanced_period'].query('post == @v').iloc[0] for v in [0, 1]]
N = lambda x: f'{float(x):,.0f}'
F = lambda x, n=1: f'{float(x):.{n}f}'.replace('-', '−')
P = lambda x: '< 0.001' if x < .001 else f'{x:.3f}'
EQP = lambda x: '< 0.001' if x < .001 else '= ' + P(x)
DATE = lambda x: pd.Timestamp(x).strftime('%B %Y')
def row(name, **filters):
    frame = D[name]
    for k, v in filters.items(): frame = frame[frame[k] == v]
    assert len(frame) == 1, (name, filters, len(frame))
    return frame.iloc[0]

def table(title, headers, rows, note):
    escape = lambda x: str(x).replace('|', r'\|').replace('\n', ' ')
    lines = ['**' + title + '**', '', '| ' + ' | '.join(headers) + ' |',
             '| ' + ' | '.join(['---'] + ['---:']*(len(headers)-1)) + ' |']
    lines += ['| ' + ' | '.join(escape(c) for c in r) + ' |' for r in rows]
    return '\n'.join(lines) + '\n\n*Note:* ' + note + '\n'

def figure(title, filename, note):
    return f'**{title}**\n\n![](assets/{filename}){{width=6.35in}}\n\n*Note:* {note}\n'

# Invariants and narrative guards. These do not certify construct validity.
assert S['n'] == D['platforms']['total'].sum() == D['monthly_all']['n'].sum()
assert S['named_n'] == D['outlets']['total'].sum()
assert S['months'] == pre['months'] + post['months']
assert post['threat'] > pre['threat'] and post['SKILLS'] < pre['SKILLS']
assert post['skills_n'] > pre['skills_n']
assert ba['threat'] > bp['threat'] and ba['skills'] < bp['skills']
assert np.allclose(D['balanced_period'][['risk_only','both','skills_only','neither']].sum(axis=1), 100)
assert np.allclose(D['balanced_period']['risk_only'] + D['balanced_period']['both'], D['balanced_period']['threat'])
assert np.allclose(D['balanced_period']['skills_only'] + D['balanced_period']['both'], D['balanced_period']['skills'])
assert row('outlet_models', model='Outlet and month effects')['estimate'] < 0
assert D['duplicates'].eval('threat_match < threat_unmatch').all()

plt.rcParams.update({'font.family':'serif','font.serif':['Times New Roman'],
    'font.size':10, 'axes.spines.top':False,'axes.spines.right':False,
    'axes.labelsize':10,'xtick.labelsize':9,'ytick.labelsize':9,
    'savefig.dpi':240, 'axes.grid':False})
launch = pd.Timestamp('2022-12-01')
def timeaxis(ax, markers=True):
    ax.xaxis.set_major_locator(mdates.YearLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y'))
    ax.grid(axis='y', color='.90', lw=.6)
    if markers:
        ax.axvline(launch,color='black',ls='--',lw=.9)
    ax.set_xlim(pd.Timestamp(S['start']), pd.Timestamp(S['end']))

m = D['monthly_all']; mw = D['monthly_web']
fig, ax = plt.subplots(figsize=(7.1,3.75))
ax.plot(pd.to_datetime(m.ym),m.n,color='black',lw=1.5,label='All platforms')
ax.plot(pd.to_datetime(mw.ym),mw.n,color='.5',ls='-',lw=1.4,label='Web only')
ax.set_ylabel('Retrieved items per month'); timeaxis(ax)
ax.legend(frameon=False,ncol=1,loc='upper left')
ax.text(launch,.99,' Launch',transform=ax.get_xaxis_transform(),va='top',fontsize=9)
for when,label,tx,ty in [('2024-03-13','Parliament AI Act vote\nMarch 2024','2023-05-01',4000),
 ('2024-08-01','AI Act in force\nAugust 2024','2024-04-01',800)]:
    date=pd.Timestamp(when)
    yy=np.interp(mdates.date2num(date),mdates.date2num(pd.to_datetime(m.ym)),m.n)
    ax.annotate(label,xy=(date,yy),xytext=(pd.Timestamp(tx),ty),fontsize=8.5,
       arrowprops={'arrowstyle':'-','color':'.45','lw':.65},color='.25',ha='left')
fig.tight_layout(); fig.savefig(ASSETS/'volume.png'); plt.close(fig)

fig, axes = plt.subplots(1,2,figsize=(7.1,3.6))
for ax, frame, title, skill in [(axes[0],m,'All retrieved items','SKILLS'),
 (axes[1],D['monthly_balanced'],'Fixed outlets, equal weights','skills')]:
    for col,label,color,style in [('threat','Threat','black','-'),('opportunity','Opportunity','.45','--'),(skill,'Skills','.25',':')]:
        ax.plot(pd.to_datetime(frame.ym),frame[col],color=color,ls=style,lw=1.25,label=label)
    ax.set_title(title,fontsize=11); ax.set_ylim(0,50); timeaxis(ax)
axes[0].set_ylabel('Share of items (%)')
handles, labels = axes[0].get_legend_handles_labels()
fig.legend(handles,labels,loc='lower center',ncol=3,frameon=False,bbox_to_anchor=(.5,-.03))
fig.tight_layout(rect=(0,.05,1,1)); fig.savefig(ASSETS/'shares.png',bbox_inches='tight'); plt.close(fig)

# Plot platform series from the scored corpus.
platform_month = A/'platform_monthly.csv'
if not platform_month.exists():
    raise RuntimeError('Run analysis script with platform_monthly.csv export before building.')
pm = pd.read_csv(platform_month)
platforms = D['platforms'].SOURCE_TYPE.tolist()
fig, axes = plt.subplots(4,3,figsize=(7.1,8.0),sharex=True)
for ax, platform in zip(axes.flat,platforms):
    z=pm[pm.SOURCE_TYPE==platform].copy(); z.ym=pd.to_datetime(z.ym)
    grid=pd.date_range(z.ym.min(),pd.Timestamp(S['end']).replace(day=1),freq='MS')
    z=z.set_index('ym').reindex(grid).fillna({'n':0})
    ax.plot(z.index,z.n,color='black',lw=1.05)
    ax.set_title(platform,fontsize=10); timeaxis(ax)
    ax.tick_params(axis='x',labelsize=8); ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.set_ylim(bottom=0)
fig.supylabel('Retrieved items per month',fontsize=10)
fig.tight_layout(); fig.savefig(ASSETS/'platforms.png'); plt.close(fig)

V={}
for k,v in {'N':S['n'],'MONTHS':S['months'],'PRE_MONTHS':pre['months'],'POST_MONTHS':post['months'],
 'NAMED_N':S['named_n'],'NAMED_OUTLETS':S['named_outlets'],'BALANCED_OUTLETS':S['balanced_outlets'],
 'BALANCED_N':S['balanced_n'],'SUSPECT_N':S['suspect_n'],'FB_N':S['fb_n'],'FB_MATCHED':S['fb_matched'],
 'DID_OUTLETS':S['did_outlets'],'DID_TABLOIDS':S['did_tabloids']}.items(): V[k]=N(v)
V.update(START=DATE(S['start']),END=DATE(S['end']),WEB_SHARE=F(100*S['web_n']/S['n']),
 NAMED_SHARE=F(100*S['named_n']/S['n']),NAMED_WEB_SHARE=F(100*S['named_n']/S['web_n']),
 FB_MATCH_SHARE=F(100*S['fb_matched']/S['fb_n']),
 FORUM_SHARE=F(100*row('forums',FROM='forum.hr')['n']/D['forums']['n'].sum()),
 INSTAGRAM_N=N(row('platforms',SOURCE_TYPE='instagram').total),INSTAGRAM_FIRST=row('platforms',SOURCE_TYPE='instagram')['first'])
for key, col in [('TH','threat'),('SK','SKILLS'),('OP','opportunity'),('SK_N','skills_n')]:
    V[key+'_PRE']=F(pre[col]); V[key+'_POST']=F(post[col])
V['SK_CHANGE_ABS']=F(pre['SKILLS']-post['SKILLS'])
for key,col in [('BAL_TH','threat'),('BAL_SK','skills')]:
    V[key+'_PRE']=F(bp[col]); V[key+'_POST']=F(ba[col])
for key,ind in [('TH','threat'),('SK','SKILLS')]:
    r=row('common_trend',indicator=ind);V[key+'_ADJ']=F(r.estimate,2)
    V[key+'_ADJ_P']= EQP(r.p) if key=='SK' else P(r.p)
for key,sample,term in [('LEVEL','all','post'),('SLOPE','all','ramp'),('WEB_LEVEL','web','post')]:
    r=row('volume_its',sample=sample,term=term);V[key]=F(r.estimate,1)
    V[key+'_P']=EQP(r.p) if key=='SLOPE' else P(r.p)
for i,r in D['breakpoints'].iterrows():
    V[f'BREAK_{i+1}']=DATE(r.point);V[f'BREAK_{i+1}_LO']=DATE(r.lower);V[f'BREAK_{i+1}_HI']=DATE(r.upper)
V['PLACEBO_LEVEL_P']=P(row('placebo',term='placebo').p)
V['PLACEBO_SLOPE_P']=P(row('placebo',term='pramp').p)
did=row('outlet_models',model='Outlet and month effects')
V.update(DID_PP=F(did.estimate*100,2),DID_SE_PP=F(did.se*100,2),DID_P=P(did.p),
 DID_EQUAL_PP=F(row('outlet_models',model='Equal outlet-month weights').estimate*100,2),
 LOO_MIN_PP=F(D['leave_one_out'].estimate.min()*100,2),LOO_MAX_PP=F(D['leave_one_out'].estimate.max()*100,2),
 LOO_MIN_P=P(D['leave_one_out'].p.min()),LOO_MAX_P=P(D['leave_one_out'].p.max()),
 CLEAN_TH=F(row('sensitivity',sample='Excluding suspect dates').estimate,2))
for w in [18,24]:
    r=row('windows',window=w); V['W'+str(w)]=F(r.estimate,2);V[f'W{w}_P']=P(r.p)
w24=row('windows',window=24);V.update(W24_PRE=N(w24.pre_n),W24_POST=N(w24.post_n),W24_HOLM=P(w24.holm_p))
for v,suffix in [(0,'PRE'),(1,'POST')]:
    r=row('duplicates',post=v)
    V['DUP_'+suffix]=F(r.share);V['DUP_TH_'+suffix]=F(r.threat_match);V['UNMATCH_TH_'+suffix]=F(r.threat_unmatch)
e=row('engagement',model='additive',term='threatTRUE')
V.update(ENG_BETA=F(e.estimate,3),ENG_SE=F(e.se,3),ENG_OP_BETA=F(row('engagement',model='additive',term='opportunityTRUE').estimate,3))
for k,model in [('POS','positive'),('PPML','ppml')]:
    r=row('engagement',model=model,term='threatTRUE');V['ENG_'+k+'_BETA']=F(r.estimate,3);V['ENG_'+k+'_P']=P(r.p)
V['ZERO_PRE']=F(D['engagement_year'].iloc[0].zero_pct);V['ZERO_LAST']=F(D['engagement_year'].iloc[-1].zero_pct)

V['TABLE_PLATFORMS']=table('Table 1. Platform composition', ['Platform','Pre','Post','Total','Share (%)','First item'],
 [[r.SOURCE_TYPE,N(r.pre),N(r.post),N(r.total),F(r.share),r['first']] for _,r in D['platforms'].iterrows()],
 'Pre: January 2021–November 2022. Post: December 2022–August 2026. Shares use all retrieved items. First item is the first qualifying observation for that platform.')
outlets=D['outlets'].sort_values(['outlet_type','outlet'])
V['TABLE_OUTLETS']=table('Table 2. Named web outlets', ['Outlet','Type','Pre','Post','Total','Months'],
 [[r.outlet,r.outlet_type,N(r.pre),N(r.post),N(r.total),N(r.months)] for _,r in outlets.iterrows()],
 'Authors’ domain-based categories. Months counts months with at least one qualifying item. Forbes is configured but contributes no items. Counts refer to web items only.')
vol=D['volume_its'];lab={'t':'Pre-period slope','post':'Level at launch','ramp':'Slope change'}
V['TABLE_VOLUME']=table('Table 3. Segmented monthly volume models', ['Sample / coefficient','Estimate','SE','p'],
 [[('All: ' if r['sample']=='all' else 'Web: ')+lab[r.term],F(r.estimate,2),F(r.se,2),P(r.p)] for _,r in vol.iterrows()],
 'Outcome: monthly item count. Level: items; slopes: items per month. Models include intercept, linear time, post and centred post-by-time interaction. Andrews HAC errors; '+N(S['months'])+' months per model. Residual autocorrelation remains: pooled Breusch–Godfrey test, order three, p '+EQP(vol.iloc[0].bg_p)+'.')
labels={'threat':'Threat composite','opportunity':'Opportunity composite','JOB_LOSS':'Job loss','JOB_CREATION':'Job creation',
 'TRANSFORMATION':'Transformation','SKILLS':'Skills','REGULATION':'Regulation','PRODUCTIVITY':'Productivity','INEQUALITY':'Inequality','FEAR_RESISTANCE':'Fear / resistance'}
V['TABLE_INDICATORS']=table('Table 4. Indicator prevalence and common-trend estimates', ['Indicator','Pre (%)','Post (%)','Post coef.','SE','p'],
 [[labels[r.indicator],F(pre[r.indicator]),F(post[r.indicator]),F(r.estimate,2),F(r.se,2),P(r.p)] for _,r in D['common_trend'].iterrows()],
 'Pre/post values are equally weighted monthly means. Coefficients and standard errors are percentage points from separate share ~ time + post regressions with Andrews HAC errors. Categories overlap; probabilities are unadjusted exploratory diagnostics.')
cats={'risk_only':'Threat without skills','both':'Threat and skills','skills_only':'Skills without threat','neither':'Neither'}
V['TABLE_JOINT']=table('Table 5. Risk and skills within a fixed outlet set', ['Item category','Pre (%)','Post (%)','Change (pp)'],
 [[v,F(bp[k]),F(ba[k]),F(ba[k]-bp[k])] for k,v in cats.items()],
 N(S['balanced_outlets'])+' named outlets with qualifying items in every month; equal outlet weights within month, then equal month weights within period. Categories are mutually exclusive and exhaustive. Percentages may differ slightly from one hundred because of rounding.')
t=D['types']; tr=[]
for ty in ['Tabloid','Quality','Regional','Public','Tech','Business','Other web','Non-web']:
    b=row('types',type=ty,post=0);a=row('types',type=ty,post=1)
    tr.append([ty,N(b['n']+a['n']),F(b.threat),F(a.threat),F(b.skills),F(a.skills)])
V['TABLE_TYPES']=table('Table 6. Counts and indicators by outlet type',['Type','Items','Threat pre','Threat post','Skills pre','Skills post'],tr,
 'Shares are item-weighted percentages within type and period. Named types include web items only. Other web and non-web are residual groups. “Quality” denotes the authors’ comparison category, not an official rating.')
V['TABLE_WINDOWS']=table('Table 7. Window sensitivity for the threat share',['Requested ±mo.','Pre / post','Estimate','SE','p','Holm p'],
 [[N(r.window),N(r.pre_n)+' / '+N(r.post_n),F(r.estimate,2),F(r.se,2),P(r.p),P(r.holm_p)] for _,r in D['windows'].iterrows()],
 'Post coefficient in share ~ time + post; percentage points; Andrews HAC errors. Requested windows are truncated by data availability. Holm adjustment covers these '+N(len(D['windows']))+' tests. Newey–West lag-three probabilities are '+P(row('windows',window=18).nw_p)+' for the eighteen-month window and '+P(row('windows',window=24).nw_p)+' for the twenty-four-month window.')
V['TABLE_SENSITIVITY']=table('Table 8. Sample and measurement sensitivity',['Specification','Estimate (pp)','SE','p'],
 [[r['sample'],F(r.estimate,2),F(r.se,2),P(r.p)] for _,r in D['sensitivity'].iterrows()],
 'Full-period common-trend regressions. Balanced specification averages within-outlet shares equally across continuously represented outlets. Threat without fear is job loss or inequality. It changes the construct as well as the estimate.')
englabels={'additive':'Log(1+y), source + month','source_month':'Log(1+y), source × month','positive':'Log(y), positive only','ppml':'Poisson, source × month'}
er=[]
for key,label in englabels.items():
    a=row('engagement',model=key,term='threatTRUE');b=row('engagement',model=key,term='opportunityTRUE')
    er.append([label,F(a.estimate,3)+' ('+F(a.se,3)+')',P(a.p),F(b.estimate,3)+' ('+F(b.se,3)+')',P(b.p),N(a['n'])])
V['TABLE_ENGAGEMENT']=table('Table 9. Engagement specification sensitivity',['Model','Threat (SE)','p','Opportunity (SE)','p','N'],er,
 'Web items. Coefficients are on the stated outcome/link scale, not percentage effects from log(1+y). Source-clustered standard errors in parentheses. Positive-only model has additive source and month effects. Retained samples differ with singleton and all-zero-group removal; all models include both indicators.')
V['TABLE_DID']=table('Table C2. Outlet comparison specifications',['Specification','Estimate (pp)','SE','p','N'],
 [[r.model,F(100*r.estimate,2),F(100*r.se,2),P(r.p),N(r['n'])] for _,r in D['outlet_models'].iterrows()],
 'Post × tabloid coefficient; two-way outlet and calendar-month clustering. First row reproduces the earlier group-plus-month specification; second adds outlet effects. Final row averages items within outlet-month and weights observed cells equally. Small-cluster inference remains a limitation.')
V['TABLE_LOO']=table('Table C3. Leave-one-out outlet comparisons',['Excluded outlet','Estimate (pp)','SE','p'],
 [[r.excluded,F(100*r.estimate,2),F(100*r.se,2),P(r.p)] for _,r in D['leave_one_out'].iterrows()],
 'Outlet-and-month fixed effects; two-way clustering by outlet and month. Each row removes the named outlet and re-estimates the same interaction.')
V['TABLE_ZEROS']=table('Table C4. Recorded web interactions',['Year','Web items','Zero','Zero (%)','Positive'],
 [[N(r.yr).replace(',',''),N(r['n']),N(r.zero),F(r.zero_pct),N(r.positive)] for _,r in D['engagement_year'].iterrows()],
 'Final year covers January–August. No distinction between structural zero and an unavailable metric coded as zero has been validated. This table describes recorded values.')
V['TABLE_POST_SLOPES']=table('Table C5. Post-period indicator trends',['Indicator','Slope (pp/mo.)','SE','p'],
 [[labels[r.indicator],F(r.estimate,3),F(r.se,3),P(r.p)] for _,r in D['post_slopes'].iterrows()],
 'Separate monthly share ~ time regressions over December 2022–August 2026, with Andrews HAC errors. Slopes are percentage points per month; no causal interpretation or multiple-testing adjustment is imposed.')
date_rows=[]
for _,r in D['suspect_dates'].iterrows():
    if '8782' in r.URL: finding='FER event page: April 2023 programme'
    elif '9051' in r.URL: finding='Pula event page: April 2023 programme'
    elif 'osiguranje' in r.URL: finding='Publisher listing: 12 September 2023'
    else: finding='Publisher page: 3 March 2023'
    link='https://www.osiguranje.hr/Default.aspx?vrsta=&zadnji=22360' if 'osiguranje' in r.URL else r.URL
    date_rows.append([pd.Timestamp(r.DATE).strftime('%d %b %Y'),f'[{r.FROM}]({link})',finding])
V['TABLE_DATES']=table('Table C1. Flagged chronology inconsistencies',['Stored date','Source link','Publisher evidence'],date_rows,
 'Targeted checks of stored pre-release ChatGPT mentions. Event timing is evidence of inconsistency, not a substitute publication timestamp. All flagged records are excluded together in Table 8.')
dict_text=[]
for key in labels:
    z=D['dictionary'][D['dictionary'].indicator==key]
    if z.empty: continue
    dict_text += ['**'+labels[key]+'.** '+str(z.iloc[0].description)+'. Expressions: '+ '; '.join(z.keyword)+'.']
V['DICTIONARY_APPENDIX']='\n\n'.join(dict_text)
V['FIG_VOLUME']=figure('Figure 1. Monthly coverage volume','volume.png',
 'Retrieved counts. Dashed line: December 2022, the first full post-launch month. AI Act dates are contextual annotations, not estimated effects.')
V['FIG_SHARES']=figure('Figure 2. Monthly lexical indicators','shares.png',
 'Unsmoothed monthly shares. Right panel: '+N(S['balanced_outlets'])+' named outlets with items in every month, equally weighted. Dashed vertical line marks December 2022. Indicators are non-exclusive and unvalidated.')
V['FIG_PLATFORMS']=figure('Figure C1. Coverage by platform','platforms.png',
 'Separate vertical scales. Dashed line: December 2022. Each platform panel begins with its first qualifying item. Instagram and other late-entry platforms have no pre-launch observations.')

template=(ROOT/'R/04_working_paper_template.md').read_text(encoding='utf-8')
used=set(re.findall(r'\{\{([A-Z0-9_]+)\}\}',template));assert used <= V.keys(),used-V.keys()
body=re.sub(r'\{\{([A-Z0-9_]+)\}\}',lambda m:V[m[1]],template)
assert '{{' not in body
abstract=(f'We describe AI-and-labour coverage in a Croatian media-monitoring corpus of {V["N"]} items from {V["START"]} to {V["END"]}. '
 f'Using monthly time-series summaries, dictionary indicators and outlet comparisons, we distinguish growth in recorded coverage from changes in its composition. '
 f'The pooled count series accelerates after ChatGPT’s release. The mean monthly threat-indicator share rises from {V["TH_PRE"]}% to {V["TH_POST"]}%, '
 f'while the skills share falls from {V["SK_PRE"]}% to {V["SK_POST"]}%, despite an increase in skills-related item counts. A fixed-outlet, equal-weight analysis shows the same directional contrast. '
 'The tabloid differential is negative but rests on few outlet clusters. Engagement associations depend on outcome specification and the treatment of zeros. '
 'The evidence is descriptive: there is no untreated comparison, and suspect dates and dictionaries awaiting human validation limit stronger interpretations. '
 'The paper contributes a transparent account of changing labour-related narratives and a framework for validating whether risk discussion is accompanied by practical adaptation information.')
front=f'''::: {{.paper-front}}
WORKING PAPER

# After ChatGPT

Media Coverage of AI’s Labour-Market Implications in Croatia

Evidence from a monitored media corpus, 2021–2026

Petra Palić · Luka Šikić · Antea Barišić

Petra Palić — Croatian Catholic University, University Department of Sociology

Luka Šikić — Croatian Catholic University, University Department of Communication Studies

Antea Barišić — University of Zagreb, Faculty of Economics and Business

29 September 2026

**Abstract**

{abstract}

**Keywords:** artificial intelligence; labour market; media coverage; framing; Croatia

Independent working paper · Preliminary research for discussion
:::

'''
md=front+body
(OUT/(STEM+'.md')).write_text(md,encoding='utf-8')
(A/'text_values.json').write_text(json.dumps(V,ensure_ascii=False,indent=2),encoding='utf-8')
# Local bibliography copy fixes the compound family name while preserving source bibliography.
bib=(ROOT/'references.bib').read_text(encoding='utf-8').replace('Bernal, James Lopez and','{Lopez Bernal}, James and')
bib=bib.replace('de Vreese, Claes H.','{de Vreese}, Claes H.')
bib+='\n'+(ROOT/'R/04_working_paper_references.bib').read_text(encoding='utf-8')
(OUT/'references.bib').write_text(bib,encoding='utf-8')
css='''body{font-family:Georgia,"Times New Roman",serif;line-height:1.55;color:#161616;max-width:880px;margin:60px auto;padding:0 32px;font-size:17px}h1,h2,h3{color:#000;font-family:inherit;line-height:1.2}h1{font-size:1.55em;margin-top:2em}h2{font-size:1.18em;margin-top:1.6em}.paper-front{text-align:center;border-bottom:1px solid #bbb;padding-bottom:2.5em;margin-bottom:3em}.paper-front h1{font-size:2.6em;margin-top:.7em}.paper-front p:nth-of-type(2){font-size:1.5em}.paper-front p:nth-of-type(3){font-size:1.05em}.paper-front p:nth-of-type(5),.paper-front p:nth-of-type(6),.paper-front p:nth-of-type(7){font-size:.85em;line-height:1.4;margin:.35em 0}.paper-front p:nth-last-of-type(3){text-align:justify;font-size:.97em}.paper-front p:last-child{font-size:.8em;letter-spacing:.04em}table{border-collapse:collapse;width:100%;font-size:.83em;line-height:1.35;margin:1em 0;border-top:1px solid #333;border-bottom:1px solid #333}th{border-bottom:1px solid #555;font-weight:600}th,td{padding:7px 8px;vertical-align:top}tr:nth-child(even){background:#fafafa}img{max-width:100%;height:auto;display:block;margin:auto}a{color:#253d53;text-underline-offset:2px}#refs{font-size:.88em;line-height:1.4}.csl-entry{margin:.6em 0}.display.math{overflow-x:auto}code{font-size:.83em;overflow-wrap:anywhere}@media(max-width:640px){body{font-size:16px;margin:20px auto;padding:0 15px}table{font-size:.72em}th,td{padding:5px 3px}.paper-front h1{font-size:2em}}@media print{body{max-width:none;font-size:11pt;margin:0}.paper-front{page-break-after:always}h1,h2{break-after:avoid}table,img{break-inside:avoid}}
'''
(ASSETS/'paper.css').write_text(css,encoding='utf-8')
base=[str(PANDOC),STEM+'.md','--from=markdown+tex_math_dollars+fenced_divs','--standalone','--citeproc','--bibliography=references.bib','--metadata=lang:en-GB','--metadata=pagetitle:After ChatGPT: A Croatian Media Working Paper']
subprocess.run(base+['--to=html5','--mathml','--embed-resources','--css=assets/paper.css','-o',STEM+'.html'],cwd=OUT,check=True)
subprocess.run(base+['--to=docx','-o',STEM+'.docx'],cwd=OUT,check=True)

doc=Document(OUT/(STEM+'.docx'))
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.85)
sec.left_margin=sec.right_margin=Inches(.95)
sec.header_distance=sec.footer_distance=Inches(.4)
sec.different_first_page_header_footer=True
for style in doc.styles:
    if style.type in [1,2]:
        style.font.name='Times New Roman'; style.font.color.rgb=RGBColor(0,0,0)
        for rf in style.element.xpath('.//w:rFonts'):
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
                rf.attrib.pop(qn('w:'+key),None)
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.13
normal.paragraph_format.space_after=Pt(6)
for nm,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
    st=doc.styles[nm];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
for nm in ['Body Text','First Paragraph']:
    if nm in doc.styles:
        doc.styles[nm].base_style=normal;doc.styles[nm].font.size=Pt(11)
        doc.styles[nm].paragraph_format.space_after=Pt(6)
        doc.styles[nm].paragraph_format.line_spacing=1.13
front_done=False
for p in doc.paragraphs:
    txt=p.text.strip();pf=p.paragraph_format
    pf.widow_control=True
    if txt.startswith('1. Introduction'):
        front_done=True;pf.page_break_before=True
    if not front_done:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_after=Pt(9)
        if txt=='WORKING PAPER':
            pf.space_before=Pt(7)
            for r in p.runs:r.font.size=Pt(11);r.font.bold=True
        elif txt=='After ChatGPT':
            p.style=doc.styles['Title'];pf.space_before=Pt(16);pf.space_after=Pt(8)
            for r in p.runs:r.font.size=Pt(24);r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
        elif txt.startswith('Media Coverage'):
            for r in p.runs:r.font.size=Pt(16);r.font.bold=True
        elif txt.startswith('Evidence from'):
            for r in p.runs:r.font.size=Pt(11)
        elif ' — ' in txt:
            pf.space_after=Pt(3)
            for r in p.runs:r.font.size=Pt(9.5)
        elif txt==abstract:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.space_after=Pt(10)
            for r in p.runs:r.font.size=Pt(10.5)
        elif txt.startswith('Keywords:') or txt.startswith('Independent working'):
            for r in p.runs:r.font.size=Pt(9.5)
        continue
    if p.style.name not in ['Heading 1','Heading 2','Heading 3']:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if txt in ['References','Appendix A. Measurement specification','Appendix B. Human-validation protocol','Appendix C. Data-quality and sensitivity details']:
        pf.page_break_before=True
    if re.match(r'^(Table|Figure) [A-Z]?\d+\. ',txt):
        pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(5)
        for r in p.runs:r.font.size=Pt(10.5);r.font.bold=True
    if txt.startswith('Note:'):
        pf.space_after=Pt(9);pf.line_spacing=1.05
        for r in p.runs:r.font.size=Pt(9)
    if p._p.xpath('.//w:drawing'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True;pf.space_after=Pt(3)
    if p.style.name in ['Bibliography','References']:
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent=Inches(.2);pf.first_line_indent=Inches(-.2)
        pf.space_after=Pt(5);pf.line_spacing=1.05
        for r in p.runs:r.font.size=Pt(10)

# Size scholarly tables and keep rows/captions/notes attached deliberately.
for tb in doc.tables:
    tb.autofit=False
    ncol=len(tb.columns)
    first=[row.cells[0].text for row in tb.rows]
    maxfirst=max(map(len,first))
    wide=2.65 if maxfirst>35 else 2.2 if maxfirst>22 else 1.65
    total=6.6;rest=(total-wide)/(ncol-1)
    widths=[wide]+[rest]*(ncol-1)
    if ncol==3:widths=[1.2,1.25,4.15]
    for c,w in zip(tb.columns,widths):c.width=Inches(w)
    tblpr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'777777');borders.append(e)
    tblpr.append(borders)
    for ri,r in enumerate(tb.rows):
        pr=r._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(r.cells):
            c.width=Inches(widths[ci]);tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for ed,sz in [('top','55'),('bottom','55'),('left','70'),('right','70')]:
                e=OxmlElement('w:'+ed);e.set(qn('w:w'),sz);e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if ri==0:
                b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e);tcpr.append(b)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or ncol==3 else WD_ALIGN_PARAGRAPH.RIGHT
                pf=p.paragraph_format;pf.space_after=Pt(0);pf.space_before=Pt(0);pf.line_spacing=1.03
                pf.keep_with_next=(ri==0 or ri==len(tb.rows)-1 or len(tb.rows)<=18)
                for run in p.runs:run.font.size=Pt(9.5);run.font.bold=(ri==0)

header=sec.header.paragraphs[0];header.text='AFTER CHATGPT\tWORKING PAPER · SEPTEMBER 2026'
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.6),WD_ALIGN_PARAGRAPH.RIGHT)
for r in header.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='After ChatGPT: Media Coverage of AI’s Labour-Market Implications in Croatia'
doc.core_properties.author='Petra Palić; Luka Šikić; Antea Barišić'
doc.core_properties.subject='Independent working paper; descriptive media-corpus evidence'
doc.core_properties.keywords='AI; labour market; media; Croatia'
doc.save(OUT/(STEM+'.docx'))
manifest={}
for path in [ROOT/'data/raw/ai_labour_corpus_2021_2026.rds',ROOT/'config.yml',ROOT/'R/04_working_paper_analysis.R',ROOT/'R/04_working_paper_template.md',Path(__file__),OUT/(STEM+'.md')]:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda:f.read(4*1024*1024),b''):h.update(chunk)
    manifest[str(path.relative_to(ROOT))]=h.hexdigest()
(A/'sha256_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print('Built HTML, DOCX and shared manuscript. Numerical guards passed.')
print('Tables:',len(doc.tables),'Figures:',len(doc.inline_shapes),'Words:',len(md.split()))
